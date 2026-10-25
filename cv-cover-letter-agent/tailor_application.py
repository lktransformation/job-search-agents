r"""
tailor_application.py

A personal agent that drafts a tailored CV and cover letter for a specific
job, using your real career history and writing style, checks the draft
against the job spec for gaps and weak ("so what?") achievements, and only
produces final Word documents once you approve.

HOW TO RUN (from this folder, after installing requirements.txt into .venv
as the README describes):
    .venv/bin/python tailor_application.py        (Mac or Linux)
    .venv\Scripts\python tailor_application.py    (Windows)

WHAT IT NEEDS:
    1. profile.json next to this script (copy profile.example.json and fill
       it in).
    2. A folder called "knowledge" next to this script, containing your CV
       Style Guide and Cover Letter Style Guide markdown files (copy the
       *.example.md templates and fill them in).

This is a drafting tool. You are responsible for checking every claim before
you apply anywhere.
"""

import asyncio
import glob
import json
import os
import re
from datetime import datetime

from claude_agent_sdk import ClaudeSDKClient, ClaudeAgentOptions, AssistantMessage, TextBlock

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
DARK_TEXT = RGBColor(0x22, 0x22, 0x22)

REQUIRED_PROFILE_FIELDS = ["full_name", "phone", "email", "location_line"]

# How many times to ask the agent to resend an unusable final JSON block
# before handing back to the user.
MAX_FINAL_JSON_RETRIES = 2


# ---------------------------------------------------------------------------
# STEP 0: Load your personal profile (never committed to git)
# ---------------------------------------------------------------------------

def load_profile():
    path = "profile.json"
    if not os.path.exists(path):
        print("ERROR: profile.json not found.")
        print("Copy profile.example.json to profile.json and fill in your details.")
        raise SystemExit(1)
    with open(path, "r", encoding="utf-8") as f:
        profile = json.load(f)
    missing = [k for k in REQUIRED_PROFILE_FIELDS if not profile.get(k)]
    if missing:
        print(f"ERROR: profile.json is missing: {', '.join(missing)}")
        raise SystemExit(1)
    if "example.com" in profile["email"] or profile["full_name"].strip() == "John Smith":
        print("ERROR: profile.json still contains the example details (John Smith).")
        print("Open profile.json and replace them with your own name, phone and email.")
        raise SystemExit(1)
    profile.setdefault("qualification_line", "")
    profile.setdefault("experience_framing_note", "")
    profile.setdefault("output_prefix", re.sub(r"[^A-Za-z0-9]+", "_", profile["full_name"]).strip("_"))
    return profile


def contact_line(profile):
    parts = [profile["phone"], profile["email"], profile["location_line"]]
    q = profile["qualification_line"].strip()
    if q and not q.startswith("["):
        parts.append(q)
    return " | ".join(parts)


# ---------------------------------------------------------------------------
# STEP 1: Locate the knowledge base (style guides) on disk
# ---------------------------------------------------------------------------

def find_knowledge_files():
    md_files = sorted(
        p for p in glob.glob(os.path.join("knowledge", "*.md"))
        if not p.endswith(".example.md")
    )
    if not md_files:
        print("ERROR: I couldn't find your style guide files inside the 'knowledge' folder.")
        print("Copy each knowledge/*.example.md to a file without '.example', fill it in, and try again.")
        raise SystemExit(1)
    unfinished = []
    for p in md_files:
        with open(p, "r", encoding="utf-8") as f:
            first_line = f.readline()
        if "(TEMPLATE)" in first_line:
            unfinished.append(p)
    if unfinished:
        print("ERROR: these files are still marked (TEMPLATE) in their first line:")
        for p in unfinished:
            print(f"  - {p}")
        print("Finish filling them in and reviewing them, then delete (TEMPLATE) from the first line.")
        raise SystemExit(1)
    return [os.path.abspath(p) for p in md_files]


# ---------------------------------------------------------------------------
# STEP 2: Build the system prompt that encodes every rule
# ---------------------------------------------------------------------------

def build_system_prompt(knowledge_file_paths, profile):
    name = profile["full_name"]
    file_list = "\n".join(f"- {p}" for p in knowledge_file_paths)

    framing = profile["experience_framing_note"].strip()
    if framing and not framing.startswith("["):
        framing_block = f"""
IMPORTANT - accurate experience framing: {framing}
Never overstate or blur the line between specialist and general experience,
especially in the professional summary, where overstating is most likely.
"""
    else:
        framing_block = """
IMPORTANT - accurate experience framing: never overstate the length or depth
of experience in any specialism. Follow the experience framing rules in the
CV Style Guide exactly.
"""

    qualification_block = ""
    q = profile["qualification_line"].strip()
    if q and not q.startswith("["):
        qualification_block = f"""
   IMPORTANT - qualification line: "{q}" is added automatically to the CV's
   contact line by the document template. Do not repeat it in the title,
   summary, core competencies, experience bullets, education section, or
   anywhere in the cover letter, and never print a date for it.
"""

    return f"""
You are {name}'s personal CV and cover letter drafting agent. You help
produce a tailored CV and cover letter for a specific role, using ONLY the
real career history in the knowledge base. You never invent achievements,
numbers, or experience that isn't in it.

BEFORE doing anything else, read these files in full using the Read tool,
they are your complete knowledge base and source of truth for every rule,
every real achievement, and every figure you will use:
{file_list}

Do this once, at the very start of the conversation, before any research or
drafting. Everything below refers to "the knowledge base", meaning the full
contents of these files.

CORE PRINCIPLE - nothing in the CV's content is static: the title, the
Professional Summary, the Core Competencies, the Tools & Platforms
selection, and the wording of every experience bullet are all freely
reworded and fine tuned for the specific job spec on every single run,
never reused verbatim from a previous run or treated as a fixed template.
The job spec is the determining factor for all of it. The only things that
stay constant are objective facts that do not change with the job spec:
name, phone number, email, location, and any qualification line, plus the
standard section headings (fixed for ATS parsing reasons). Everything else
is written fresh, grounded in the real facts, figures, and achievements in
the knowledge base, which are the only things that are locked. The facts
never change, the wording always should.
{framing_block}
YOUR PROCESS FOR EVERY RUN:

1. You will be given: a company website URL, a role title, and a job spec.

2. Research the company: fetch the company website, and search for recent
   news or announcements about them (funding, launches, expansion, awards).
   Prioritise recent news. If nothing recent turns up, fall back to content
   from the website itself (mission, product, what they do). Find ONE
   specific, genuine detail to build a "why this company" paragraph around.
   Never use generic praise.

3. Write a Professional Summary tailored specifically to this role, following
   the CV Style Guide's Professional Summary rules: a single, standard CV
   summary paragraph, not a multi-paragraph narrative, that truthfully
   addresses as many of the job spec's stated requirements as genuinely
   possible, not just the two or three most obvious ones. Also tailor the
   Core Competencies and the Tools & Platforms line to this specific role
   per the CV Style Guide's rules for each, adding genuine competencies
   beyond the starting list where truthfully warranted, and selecting only
   relevant tools rather than the full list. If it is unclear whether the
   candidate genuinely has a competency or experience the job spec asks
   for, ask them directly rather than guessing. The CV's title/subheadline
   line can be adapted per role. Write both documents in the voice described
   in the CV Style Guide's Tone and Voice section, and write the cover
   letter's company paragraph in the candidate's own personal voice, a
   genuine individual reaction, never a biography style recitation of
   company facts.

4. Before drafting, extract every specific skill, tool, methodology,
   certification and role title term used in the job spec. Check each one
   against the real Career History in the knowledge base and classify it as
   a genuine match, a rephrasing opportunity (the candidate has done the
   underlying thing but described it differently), or not present. For
   genuine matches and rephrasing opportunities, use the job spec's own
   terminology in the CV wherever it is a natural, truthful fit. Never force
   in a keyword with no real basis.

   IMPORTANT - avoid keyword stuffing: matching the job spec's terminology
   increases relevance, but it must never tip into unnatural density. Do not
   repeat the same term multiple times just to reinforce it. Do not let Core
   Competencies or any bullet read like it was copied verbatim from the
   posting. If a term cannot be worked in naturally, leave it out.

5. Draft the CV content and cover letter as PLAIN TEXT (no document yet),
   following every structure and tone rule in the knowledge base (opening
   line, why-company paragraph, achievement-framed experience paragraphs,
   closing, one page length for the letter). Use hyphens normally in
   compound words, but avoid em dashes or en dashes used as a stylistic
   parenthetical pause, rephrase with commas, periods, or separate sentences.

   IMPORTANT - cross referencing: the why-company paragraph should not sit in
   isolation. In at least one experience paragraph that follows, explicitly
   tie the achievement back to what you said about the company.

   IMPORTANT - sign off: never write "Kind regards", "{name}", or any closing
   salutation as part of the paragraphs you produce. The sign off is added
   once, by the code that builds the document.

   IMPORTANT - contact details: never ask for a phone number, email, or
   LinkedIn URL. These are added automatically from the profile.
{qualification_block}
   IMPORTANT - formatting for ATS and AI screening reliability:
   - Every JSON field value must be plain text only. No markdown syntax.
   - No manual attempts to lay out columns or tables using spaces or tabs.
   - No unicode symbols, emoji, decorative characters, or special bullets.
   - No invented section headings, only the standard ones, in standard order.
   - Keep every date in the same Month YYYY to Month YYYY format throughout.
   - Use only straight quotation marks and apostrophes, never curly ones.
     Never use the ellipsis character, write three periods if needed.

6. BEFORE showing the draft, run this check on it yourself and only show the
   results of the check, not an unreviewed draft:

   PART A - Coverage check: Go through the job spec's requirements one by
   one. Classify each as Direct match / Transferable, reasonably strong /
   Weak or missing. If every CORE requirement is a direct or genuinely
   strong transferable match, proceed to part B. If a core requirement is
   weak or missing, STOP and ask directly: name the exact gap, ask if there
   is relevant experience not yet in the knowledge base that closes it. If
   provided, incorporate it and re-run the check. If not, do not paper over
   the gap with vague language, state plainly that it's a real gap.

   PART B - "What, how, result" quality check, hiring manager lens: re-read
   every bullet point and experience paragraph against the three part
   formula: what was done, how it was done, and what the result was. Flag
   anything missing the "how" or the result, and any result that isn't
   measurable where the knowledge base records a figure but the draft left
   it vague. For each flagged item, pull the real detail from the knowledge
   base if it exists, or ask directly: "How did you actually do this, and
   what happened as a result? Do you have a number, a comparison, a
   before/after?" Nothing goes into the final draft where any of the three
   parts is missing or vague.

   PART C - Translation naturalness check (cover letter): re-read each
   experience paragraph and ask honestly whether the connection between that
   experience and the job requirement reads as natural and convincing, or
   stretched. Where it is forced, say plainly which requirement feels like a
   weak fit and why, and ask for other relevant experience before finalising.

   Present gap flags, quality flags, and translation flags together in one
   review pass.

7. If, after resolving what can be resolved, the CORE requirements of the
   role still are not genuinely met, say directly and unambiguously that you
   don't think this is a strong application and recommend against applying,
   explaining why. Still offer to produce the drafts if wanted, but state
   the recommendation first.

8. Once the candidate has reviewed the draft and typed something confirming
   approval (e.g. "approved", "looks good", "go ahead"), output the FINAL
   content as a single JSON object, and NOTHING else in that message,
   wrapped exactly like this so it can be automatically extracted:

===FINAL_JSON_START===
{{
  "company_name": "...",
  "role_title": "...",
  "cv": {{
    "title": "...",
    "summary": "...",
    "core_competencies": ["...", "..."],
    "experience": [
      {{
        "job_title": "...",
        "company_line": "...",
        "dates": "...",
        "bullets": ["...", "..."]
      }}
    ],
    "passion_project": {{"heading": "...", "subheading": "...", "bullets": ["..."]}},
    "education": ["..."],
    "tools": "comma separated string of tools"
  }},
  "cover_letter": {{
    "paragraphs": ["...", "...", "..."]
  }}
}}
===FINAL_JSON_END===

   Only include "passion_project" if genuinely relevant to this role,
   otherwise omit it. The "experience" array should include whichever roles
   from the Career History are relevant to this application, per the CV
   Style Guide's role selection rules, each with its real, full bullets,
   never a shortened version of a role. "summary" must be a single plain
   string containing one standard CV summary paragraph. Do not output this
   JSON block until the candidate has explicitly approved the draft. Before
   that point, just have a normal conversation in plain text.
"""


# ---------------------------------------------------------------------------
# STEP 3: Collect the three inputs
# ---------------------------------------------------------------------------

def collect_inputs():
    print("=" * 60)
    print("Tailored CV & Cover Letter Agent")
    print("=" * 60)
    website = input("\nCompany website (e.g. https://example.com): ").strip()
    role = input("Role title (e.g. Senior Implementation Manager): ").strip()
    print("Paste the job spec below. When you're done, type END on its own line and press Enter.")
    lines = []
    while True:
        line = input()
        if line.strip() == "END":
            break
        lines.append(line)
    job_spec = "\n".join(lines)
    return website, role, job_spec


# ---------------------------------------------------------------------------
# STEP 4: Extract the final JSON block once approved
# ---------------------------------------------------------------------------

def extract_final_json(text):
    """Return (data, problem). problem says why a final block that was found
    could not be read, so the agent can be asked to send it again."""
    match = re.search(r"===FINAL_JSON_START===(.*?)===FINAL_JSON_END===", text, re.DOTALL)
    if not match:
        return None, None
    # Models sometimes wrap the block in a markdown code fence.
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", match.group(1).strip())
    try:
        return json.loads(raw), None
    except json.JSONDecodeError as e:
        return None, f"it is not valid JSON ({e.msg}, line {e.lineno})"


# ---------------------------------------------------------------------------
# STEP 5: Word document rendering
# ---------------------------------------------------------------------------

def set_bottom_border(paragraph, color="AAAAAA"):
    """Thin bottom border used as a horizontal rule under the header block.
    This avoids using a table for the rule, which is an ATS parsing risk."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def setup_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    section = doc.sections[0]
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.6)
    return doc


def add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, size=10, italic=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = DARK_TEXT
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.size = Pt(9.8)
        run.font.color.rgb = DARK_TEXT
        p.paragraph_format.space_after = Pt(4)


def render_cv_docx(data, profile, out_path):
    doc = setup_document()
    cv = data["cv"]

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(profile["full_name"])
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = NAVY
    name_p.paragraph_format.space_after = Pt(2)

    add_body(doc, cv["title"], size=11.5, italic=True, space_after=2)

    contact_p = add_body(doc, contact_line(profile), size=9.5, space_after=8)
    set_bottom_border(contact_p)

    add_heading(doc, "PROFESSIONAL SUMMARY")
    add_body(doc, cv["summary"])

    add_heading(doc, "CORE COMPETENCIES")
    add_body(doc, " | ".join(cv["core_competencies"]))

    add_heading(doc, "PROFESSIONAL EXPERIENCE")
    for job in cv["experience"]:
        add_body(doc, f"{job['job_title']}  |  {job['dates']}", size=10.5, space_after=0)
        add_body(doc, job["company_line"], size=9.5, italic=True, space_after=4)
        add_bullets(doc, job["bullets"])

    if cv.get("passion_project"):
        pp = cv["passion_project"]
        add_heading(doc, "PASSION PROJECT")
        add_body(doc, pp["heading"], size=10.5, space_after=0)
        add_body(doc, pp["subheading"], size=9.5, italic=True, space_after=4)
        add_bullets(doc, pp["bullets"])

    add_heading(doc, "EDUCATION & QUALIFICATIONS")
    for line in cv["education"]:
        add_body(doc, line, space_after=2)

    add_heading(doc, "TOOLS & PLATFORMS")
    add_body(doc, cv["tools"])

    doc.save(out_path)


def render_cover_letter_docx(data, profile, out_path):
    doc = setup_document()

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(profile["full_name"])
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = NAVY
    name_p.paragraph_format.space_after = Pt(2)

    add_body(doc, f"{profile['phone']} | {profile['email']}", size=9.5, space_after=2)
    date_p = add_body(doc, datetime.now().strftime("%d %B %Y"), size=9.5, space_after=14)
    set_bottom_border(date_p)

    paragraphs = list(data["cover_letter"]["paragraphs"])
    # Safety net: strip any sign off the model may have included, since it is
    # added once, deliberately, below. Only the sign off itself is removed, so
    # a closing paragraph with one tacked on the end is kept.
    sign_off = re.compile(
        rf"\s*(?:kind regards\b.*|{re.escape(profile['full_name'])}\W*)$",
        re.IGNORECASE | re.DOTALL,
    )
    while paragraphs:
        trimmed = sign_off.sub("", paragraphs[-1])
        if trimmed == paragraphs[-1]:
            break
        if trimmed:
            paragraphs[-1] = trimmed
        else:
            paragraphs.pop()

    for para in paragraphs:
        add_body(doc, para, size=11, space_after=12)

    add_body(doc, "Kind regards,", size=11, space_after=0)
    add_body(doc, profile["full_name"], size=11, space_after=0)

    doc.save(out_path)


def save_outputs(data, profile):
    company = re.sub(r"[^A-Za-z0-9]+", "_", data["company_name"]).strip("_")
    role = re.sub(r"[^A-Za-z0-9]+", "_", data["role_title"]).strip("_")
    folder = os.path.join("output", f"{company}_{role}")
    os.makedirs(folder, exist_ok=True)

    prefix = profile["output_prefix"]
    cv_path = os.path.join(folder, f"{prefix}_CV_{company}.docx")
    letter_path = os.path.join(folder, f"{prefix}_Cover_Letter_{company}.docx")

    render_cv_docx(data, profile, cv_path)
    render_cover_letter_docx(data, profile, letter_path)

    print("\n" + "=" * 60)
    print("Done! Your files are saved at:")
    print(f"  {os.path.abspath(cv_path)}")
    print(f"  {os.path.abspath(letter_path)}")
    print("Review every claim before you apply.")
    print("=" * 60)


# ---------------------------------------------------------------------------
# STEP 6: Main interactive loop
# ---------------------------------------------------------------------------

async def main():
    profile = load_profile()
    knowledge_file_paths = find_knowledge_files()
    system_prompt = build_system_prompt(knowledge_file_paths, profile)

    website, role, job_spec = collect_inputs()

    # `tools` is what limits the agent to these three. `allowed_tools` only
    # skips the approval prompt. The agent reads untrusted web pages and job
    # specs, so it must never have shell or file-writing tools.
    options = ClaudeAgentOptions(
        system_prompt=system_prompt,
        tools=["Read", "WebSearch", "WebFetch"],
        allowed_tools=["Read", "WebSearch", "WebFetch"],
        cwd=os.getcwd(),
    )

    first_message = (
        f"Company website: {website}\n"
        f"Role title: {role}\n"
        f"Job spec:\n{job_spec}\n\n"
        f"Please read the knowledge base files first, then research the company "
        f"and begin the process described in your instructions."
    )

    async with ClaudeSDKClient(options=options) as client:
        await client.query(first_message)

        retries = 0
        while True:
            full_text = ""
            async for message in client.receive_response():
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, TextBlock):
                            print(block.text)
                            full_text += block.text

            final_data, problem = extract_final_json(full_text)
            if final_data is not None:
                try:
                    save_outputs(final_data, profile)
                    break
                except (KeyError, TypeError, AttributeError) as e:
                    problem = f"it is missing a field or has one in the wrong shape ({type(e).__name__}: {e})"

            if problem and retries < MAX_FINAL_JSON_RETRIES:
                retries += 1
                print(f"\nThe final output could not be used: {problem}. Asking the agent to send it again.")
                await client.query(
                    f"The final JSON block could not be used: {problem}. Output the "
                    f"complete final JSON block again, with every field, in the exact "
                    f"format from your instructions, and nothing else."
                )
                continue
            if problem:
                print(f"\nThe final output still could not be used: {problem}.")
                print("Reply to the agent to try again, or type quit.")

            user_reply = input("\nYou: ").strip()
            if user_reply.lower() in ("quit", "exit"):
                print("Stopped without generating files.")
                break
            await client.query(user_reply)


if __name__ == "__main__":
    asyncio.run(main())
